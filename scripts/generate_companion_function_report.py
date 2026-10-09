"""Export the reviewed AI Companion Markdown report to Word and a static API inventory."""
import ast
import csv
import re
from pathlib import Path
from docx import Document
from generate_project_report import configure, table, bullet

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "reports/Bao_cao_chuc_nang_AI_Companion_2026-10-07.md"
WORD = SOURCE.with_suffix(".docx")
INVENTORY = ROOT / "reports/AI_Companion_API_inventory_2026-10-07.csv"


def export_word():
    doc = Document()
    configure(doc)
    for section in doc.sections:
        for paragraph in section.footer.paragraphs:
            for run in paragraph.runs:
                if "AI BOX Companion" in run.text:
                    run.text = run.text.replace("AI BOX Companion", "AI Companion")
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in row):
                    rows.append(row)
                index += 1
            if len(rows) > 1:
                table(doc, rows[0], rows[1:])
            continue
        heading = re.match(r"^(#{1,3}) (.+)", line)
        if heading:
            level = len(heading.group(1))
            if level == 1:
                doc.add_heading(heading.group(2), 0)
            else:
                doc.add_heading(heading.group(2), level - 1)
        elif line.startswith("- "):
            bullet(doc, line[2:])
        else:
            doc.add_paragraph(line)
        index += 1
    doc.core_properties.title = "Báo cáo chức năng hiện tại — AI Companion"
    doc.core_properties.subject = "Assistant, Roleplay, CMS và vận hành"
    doc.core_properties.author = "Avibi"
    doc.save(WORD)
    return WORD


def export_inventory():
    rows = []
    for path in sorted((ROOT / "character_chat").glob("*.py")):
        tree = ast.parse(path.read_text(encoding="utf-8-sig"), filename=str(path))
        for node in ast.walk(tree):
            if not isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
                continue
            for decorator in node.decorator_list:
                if not isinstance(decorator, ast.Call) or not isinstance(decorator.func, ast.Attribute):
                    continue
                method = decorator.func.attr.upper()
                if method not in {"GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS", "HEAD"}:
                    continue
                if not decorator.args or not isinstance(decorator.args[0], ast.Constant):
                    continue
                route = decorator.args[0].value
                if isinstance(route, str):
                    rows.append([path.relative_to(ROOT).as_posix(), node.lineno, method, route, node.name])
    rows.sort(key=lambda row: (row[3], row[2], row[0]))
    with INVENTORY.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output)
        writer.writerow(["Source file", "Handler line", "HTTP method", "Path", "Handler"])
        writer.writerows(rows)
    return len(rows)


if __name__ == "__main__":
    print(export_word())
    print(f"Static API routes: {export_inventory()}")
    print(INVENTORY)
