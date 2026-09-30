from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "Bao_cao_tong_quan_AI_BOX_Companion.docx"
PURPLE = "6D4AFF"
DARK = "17131F"
LIGHT = "F1EEFA"
GRAY = "66616F"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, header in enumerate(headers):
        shade(t.rows[0].cells[i], PURPLE)
        set_cell_text(t.rows[0].cells[i], header, True, "FFFFFF", 9)
    for ridx, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, False, DARK, 8.5)
            if ridx % 2:
                shade(cells[i], "F8F7FB")
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return t


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p


def code_block(doc, text):
    p = doc.add_paragraph()
    p.style = "Code Block"
    p.add_run(text)
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_page_number(section):
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("AI BOX Companion  •  Tài liệu nội bộ  •  ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def configure(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in [("Title", 30, DARK), ("Heading 1", 20, PURPLE), ("Heading 2", 14, DARK), ("Heading 3", 11, PURPLE)]:
        style = doc.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    if "Code Block" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Cascadia Mono"
        style.font.size = Pt(8.2)
        style.font.color.rgb = RGBColor.from_string("EDE7F6")
        style.paragraph_format.left_indent = Cm(0.35)
        style.paragraph_format.right_indent = Cm(0.35)
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(8)
        p_pr = style.element.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), DARK)
        p_pr.append(shd)


def cover(doc):
    doc.add_paragraph("AI BOX", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    run = p.add_run("BÁO CÁO TỔNG QUAN\nSẢN PHẨM VÀ KỸ THUẬT")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(29)
    run.font.color.rgb = RGBColor.from_string(PURPLE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("AI BOX Companion – Nền tảng nhân vật AI đồng hành")
    run.font.name = "Aptos Display"
    run.font.size = Pt(18)
    run.bold = True
    doc.add_paragraph("\n")
    table(doc, ["Thông tin", "Giá trị"], [
        ("Phiên bản báo cáo", "1.0"),
        ("Ngày lập", datetime.now().strftime("%d/%m/%Y")),
        ("Trạng thái", "MVP hoạt động trên môi trường production"),
        ("Đối tượng", "Ban quản lý, Product, Engineering và đối tác tích hợp"),
        ("Phân loại", "Tài liệu nội bộ – không chứa API key hoặc mật khẩu"),
    ], [5, 11])
    p = doc.add_paragraph("Production: https://nguyenthanhcongloi332004--character-chat-web.modal.run/")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    doc.add_page_break()


def build():
    doc = Document()
    configure(doc)
    cover(doc)

    heading(doc, "Mục lục nội dung", 1)
    for item in [
        "1. Tóm tắt điều hành", "2. Phạm vi và trạng thái sản phẩm", "3. Trải nghiệm người dùng",
        "4. Kiến trúc tổng thể", "5. Cấu trúc mã nguồn", "6. Character Engine và Prompt Compiler",
        "7. Relationship, Emotion và Interest", "8. Memory và Background Intelligence",
        "9. Tài liệu, tri thức và truy xuất", "10. Chat 1–1, chat nhóm và streaming",
        "11. Suggestion Engine", "12. Kiểm duyệt và an toàn", "13. Mô hình dữ liệu",
        "14. API", "15. Frontend", "16. LLM, fallback và chi phí", "17. Bảo mật",
        "18. Deployment và vận hành", "19. Kiểm thử", "20. Đánh giá hiện trạng và rủi ro",
        "21. Lộ trình SaaS", "22. Hướng dẫn bàn giao", "Phụ lục A–D"
    ]:
        bullet(doc, item)
    doc.add_page_break()

    heading(doc, "1. Tóm tắt điều hành", 1)
    doc.add_paragraph(
        "AI BOX Companion là MVP B2C cho phép người dùng trò chuyện với các nhân vật AI có cá tính ổn định, "
        "quan hệ phát triển theo thời gian, ký ức, cảm xúc, tài liệu đính kèm và khả năng hỗ trợ đời sống hằng ngày. "
        "Sản phẩm kết hợp trải nghiệm giải trí (tổng tài, người yêu, bạn thân) với ba vai trò hữu ích (đồng hành, học tập, công việc)."
    )
    table(doc, ["Năng lực", "Hiện trạng", "Giá trị"], [
        ("Character Engine", "Đã triển khai", "Mỗi nhân vật có CharacterSpec, prompt, giọng nói, mục tiêu và giới hạn riêng"),
        ("Quan hệ & cảm xúc", "Đã triển khai", "Thay đổi chậm theo hội thoại nhưng không làm trôi tính cách lõi"),
        ("Memory", "Đã triển khai ở mức MVP", "Nhớ sở thích, sự kiện và chi tiết quan trọng theo phạm vi"),
        ("Tài liệu", "Đã triển khai", "PDF/Word/TXT/MD/PPTX/CSV/XLSX trở thành tri thức trong phòng chat"),
        ("Streaming UX", "Đã triển khai", "Tin nhắn người dùng xuất hiện ngay, nhân vật có trạng thái đang nhập"),
        ("Moderation", "Đã triển khai nhiều lớp", "Từ khóa + OpenAI Moderation tùy chọn + LLM kiểm tra ngữ cảnh"),
        ("Production", "Đã triển khai Modal", "Frontend và API dùng chung một endpoint công khai"),
    ], [3.5, 3, 9.5])
    doc.add_paragraph("Kết luận quản trị: MVP đã chứng minh luồng sản phẩm chính. Trước khi mở rộng thương mại cần thay cơ chế demo user bằng xác thực thật, chuyển SQLite sang PostgreSQL/pgvector, bổ sung quan sát hệ thống, quota, thanh toán và kiểm thử tải.")

    heading(doc, "2. Phạm vi và trạng thái sản phẩm", 1)
    heading(doc, "2.1 Mục tiêu", 2)
    bullet(doc, "Tạo cảm giác nhân vật sống, có bản sắc và duy trì tính cách qua hội thoại dài.")
    bullet(doc, "Tăng retention bằng ký ức, quan hệ, phản hồi cảm xúc và gợi ý tiếp tục hội thoại.")
    bullet(doc, "Kết hợp giải trí với trợ giúp học tập, công việc và đồng hành hằng ngày.")
    bullet(doc, "Cho phép tạo nhân vật từ mô tả, bản thân hoặc tài liệu/cốt truyện.")
    heading(doc, "2.2 Các phần nằm ngoài MVP hiện tại", 2)
    bullet(doc, "Chưa có TTS/voice, avatar biểu cảm động hoặc ứng dụng mobile native.")
    bullet(doc, "Chưa có authentication/authorization tiêu chuẩn SaaS, billing, subscription và admin portal.")
    bullet(doc, "Chưa dùng vector database thực; truy xuất hiện tại dùng vector từ vựng cục bộ và cosine similarity.")
    bullet(doc, "Chưa được đánh giá tải lớn; Modal đang cấu hình tối đa một container để bảo vệ SQLite.")

    heading(doc, "3. Trải nghiệm người dùng", 1)
    heading(doc, "3.1 Nhân vật mặc định", 2)
    table(doc, ["Nhân vật", "Nhóm", "Vai trò", "Cá tính bất biến", "Phát triển theo thời gian"], [
        ("Liost", "Hữu ích", "Đồng hành", "Ấm áp, tò mò, hài hước nhẹ", "Nhiều callback, câu đùa chung và thấu hiểu hơn"),
        ("Hà", "Hữu ích", "Học tập", "Vui vẻ, hòa đồng, giải thích rõ", "Hiểu trình độ, tăng thử thách và nhịp học chung"),
        ("Antony", "Hữu ích", "Công việc", "Nhanh, thực tế, dí dỏm", "Phản biện trực tiếp hơn và nhớ ngữ cảnh dự án"),
        ("Duy", "Giải trí", "Tổng tài", "Điềm tĩnh, quyết đoán, tiết chế", "Quan tâm qua hành động và chỉ mềm mỏng khi đủ tin tưởng"),
        ("Mia", "Giải trí", "Người yêu", "Ấm áp, tinh nghịch, chú ý cảm xúc", "Thêm nghi thức chung, câu đùa riêng và dỗi nhẹ có thể làm hòa"),
        ("Bảo", "Giải trí", "Bạn thân", "Nhanh trí, cà khịa, chân thành", "Tăng cách nói tắt, callback và inside joke"),
    ], [2, 2, 2.3, 4.8, 5.2])
    heading(doc, "3.2 Các hành trình chính", 2)
    for text in [
        "Chọn preset → thêm vào danh sách nhân vật → bấm Trò chuyện → tạo phòng mới → nhận lời chào đầu tiên.",
        "Tự tạo nhân vật → nhập tên/mô tả/role/scenario/ví dụ → LLM biên dịch CharacterSpec → bắt đầu chat.",
        "Tạo bản sao AI → nhập tên, mô tả và ngày sinh tùy chọn → bổ sung cung hoàng đạo nếu có → biên dịch nhân vật.",
        "Tạo từ truyện/tài liệu → tải nguồn → phát hiện nhân vật → chọn cutoff → biên dịch tính cách → xóa nguồn tạm.",
        "Chat tài liệu → đính kèm file ngay trong phòng → hiển thị thẻ file → xử lý nền → AI dùng nội dung liên quan để trả lời.",
        "Chat nhóm → chọn từ hai AI trở lên → mỗi AI chỉ thấy user và lịch sử riêng → tag tên để chỉ định người trả lời.",
    ]:
        numbered(doc, text)

    heading(doc, "4. Kiến trúc tổng thể", 1)
    code_block(doc, """Browser / React + TypeScript
        │  REST + Server-Sent Events
        ▼
FastAPI Companion API
        ├── Character Engine + Prompt Compiler
        ├── Relationship / Emotion / State Engine
        ├── Memory + Retrieval + Document Knowledge
        ├── Suggestion Engine (chạy song song)
        ├── Moderation Pipeline
        └── Background Mood / Reflection Analysis
        │
        ├── SQLite trên Modal Volume
        ├── AVIBI Gateway → GPT-4.1-mini
        ├── AVIBI Gateway → GPT-5-nano fallback/background
        └── OpenAI omni-moderation-latest (tùy chọn)
""")
    table(doc, ["Tầng", "Thành phần", "Trách nhiệm"], [
        ("Presentation", "React 19, TypeScript, Vite", "UI nhân vật, phòng chat, typing, file, suggestion và quản lý dữ liệu"),
        ("API", "FastAPI", "REST/SSE, validation, ownership, orchestration và OpenAPI"),
        ("Domain", "engine.py, prompts.py, presets.py", "Character routing, prompt assembly, state và LLM"),
        ("Intelligence", "background.py, retrieval.py", "Mood, reflection, memory extraction và context retrieval"),
        ("Safety", "moderation.py", "Kiểm tra từ khóa, ngữ cảnh và nội dung upload"),
        ("Persistence", "SQLAlchemy + SQLite", "Character, room, message, relationship, memory, document và trace"),
        ("Runtime", "Modal", "Image, secret, volume, deploy và public ASGI endpoint"),
    ], [3, 5, 8])

    heading(doc, "5. Cấu trúc mã nguồn", 1)
    table(doc, ["Đường dẫn", "Vai trò kỹ thuật"], [
        ("character_chat/main.py", "FastAPI routes, SSE, ownership, document jobs và orchestration"),
        ("character_chat/engine.py", "LLM client, routing, state, relationship, prompt assembly và suggestion"),
        ("character_chat/presets.py", "Sáu CharacterSpec mặc định và system prompt riêng"),
        ("character_chat/prompts.py", "Prompt chung cho response, suggestion và character compiler"),
        ("character_chat/schemas.py", "Pydantic contracts cho API và structured LLM output"),
        ("character_chat/db.py", "SQLAlchemy models, session và database initialization"),
        ("character_chat/retrieval.py", "Chunking, cosine lexical retrieval, memory scoring và deduplication"),
        ("character_chat/documents.py", "Trích xuất nội dung nhiều định dạng và lưu chunks"),
        ("character_chat/moderation.py", "Safety pipeline nhiều lớp"),
        ("character_chat/background.py", "Phân tích mood, event, reflection và memory sau response"),
        ("character_chat/zodiac.py", "Suy ra cung hoàng đạo và knowledge bổ trợ khi có ngày sinh"),
        ("frontend/src/main.tsx", "State và luồng UI chính"),
        ("frontend/src/chat.css, style.css", "Thiết kế và hiệu ứng chat"),
        ("modal_app.py", "Build frontend, image, volume, secrets và Modal ASGI deployment"),
        ("tests/", "Unit/API/smoke/UI timing checks"),
        ("setup_*.ps1, secret_vault.ps1", "Thiết lập secrets và vault cục bộ mã hóa"),
    ], [6, 10])

    heading(doc, "6. Character Engine và Prompt Compiler", 1)
    heading(doc, "6.1 CharacterSpec", 2)
    doc.add_paragraph("Nhân vật không được biểu diễn bằng một câu mô tả đơn. CharacterSpec là hợp đồng dữ liệu gồm identity, personality, speech, support profile, system prompt, scenario, greeting, examples, goals, behavior rules và nguồn tri thức.")
    table(doc, ["Nhóm trường", "Ví dụ", "Tác dụng"], [
        ("Identity", "role, backstory", "Neo vai trò và bối cảnh nhân vật"),
        ("Personality", "warmth, humor, assertiveness, empathy, confidence…", "Điều chỉnh khí chất và phản ứng"),
        ("Speech", "verbosity, formality, emoji, humor style", "Giữ nhịp câu và giọng nói"),
        ("Support", "primary role, strengths, method", "Chuyên môn học tập/công việc/đồng hành"),
        ("Creative card", "system_prompt, scenario, samples", "Chỉ đạo hóa thân và tạo voice anchors"),
        ("Boundaries", "goals, behavior_rules", "Mục tiêu và hành vi không được vượt qua"),
    ], [3.5, 5.5, 7])
    heading(doc, "6.2 Prompt assembly", 2)
    code_block(doc, """Global response policy
+ Character-specific system prompt
+ Identity/personality/speech tuning
+ Scenario and support profile
+ Relationship stage and numeric state
+ Current emotion/energy/mood
+ Role-specific evolution contract
+ Current intent/support mode
+ Relevant tasks, memories, reflections and events
+ Attached-document knowledge
+ Dialogue/style examples
+ Recent user + this character history
+ Post-history Voice Lock
+ Latest user message
= Final LLM request""")
    bullet(doc, "Character card có ưu tiên sáng tạo cao nhất; lịch sử chỉ giữ continuity, không được thay personality.")
    bullet(doc, "Core Character Contract cố định role, temperament, values, rhythm, humor, confidence và boundaries.")
    bullet(doc, "Post-history Voice Lock được đặt ngay trước user query để chống personality drift sau hội thoại dài.")
    bullet(doc, "Mọi lời khuyên hoặc trả lời tài liệu vẫn phải đi qua giọng nói của nhân vật.")
    heading(doc, "6.3 Nhân vật do user tạo", 2)
    doc.add_paragraph("Character Compiler dùng structured JSON để biến input thành CharacterSpec. Prompt yêu cầu tách immutable core khỏi relationship-dependent expression, bảo toàn dữ kiện user cung cấp, suy luận bảo thủ cho trường thiếu và tạo lời chào đúng giọng. JSON được validate bằng Pydantic và thử sửa tối đa ba lần khi output không hợp lệ.")

    heading(doc, "7. Relationship, Emotion và Interest", 1)
    heading(doc, "7.1 Relationship progression", 2)
    table(doc, ["Biến", "Ý nghĩa", "Cách thay đổi"], [
        ("familiarity", "Mức quen thuộc", "Tăng chậm mỗi lượt và giảm tốc khi đã cao"),
        ("trust", "Niềm tin", "Tăng khi chia sẻ cá nhân/làm hòa; giảm khi xung đột"),
        ("affinity", "Thiện cảm", "Tăng khi tương tác tích cực; giảm khi phủ định/xúc phạm"),
        ("tension", "Căng thẳng", "Tăng bởi hành vi tiêu cực; giảm dần hoặc giảm khi repair"),
        ("respect", "Tôn trọng", "Tăng qua chia sẻ và repair; giảm khi xung đột"),
    ], [3, 5, 8])
    doc.add_paragraph("Các stage: Stranger → Acquainted → Familiar → Trusted → Close → Deeply close. Tốc độ đã được cố ý làm chậm để sự thân thiết có cảm giác được xây dựng, đồng thời role_evolution_contract bảo đảm tổng tài vẫn là tổng tài, bạn thân vẫn là bạn thân và nhân vật chuyên môn không mất chức năng.")
    heading(doc, "7.2 Emotion state machine", 2)
    bullet(doc, "Phát hiện nhanh user affect: sad, tired, angry, anxious, happy, apologetic hoặc neutral.")
    bullet(doc, "State nhân vật gồm emotion, mood, intensity và energy; có quán tính thay vì đổi tức thì.")
    bullet(doc, "Nhân vật romantic có thể hurt/sulking khi có ngữ cảnh phù hợp và softening khi được dỗ.")
    bullet(doc, "Intensity giảm sau thời gian offline; energy hồi phục theo thời gian và giảm nhẹ theo độ dài lượt chat.")
    bullet(doc, "Cảm xúc chỉ thay đổi cách biểu đạt, không được ghi đè tính cách lõi hoặc tạo hành vi thao túng.")
    heading(doc, "7.3 Intent/Support mode", 2)
    doc.add_paragraph("Rule-based classifier xác định study, work, health, lifestyle hoặc companion. Mỗi mode bổ sung chỉ dẫn chuyên môn, nhưng nhân vật vẫn trả lời bằng voice riêng. Health chỉ hỗ trợ wellbeing chung, không chẩn đoán hoặc giả làm bác sĩ.")

    heading(doc, "8. Memory và Background Intelligence", 1)
    table(doc, ["Loại", "Phạm vi", "Ví dụ"], [
        ("USER_GLOBAL", "Toàn user", "Sở thích có thể dùng ở nhiều nhân vật"),
        ("CHARACTER_PRIVATE", "User + một nhân vật", "Kỷ niệm riêng trong chat 1–1"),
        ("ROOM_SHARED", "Một phòng", "Sự kiện và tri thức chung của chat nhóm"),
    ], [3.5, 4.5, 8])
    doc.add_paragraph("Memory retrieval xếp hạng theo độ tương đồng từ vựng, importance, recency, loại cảm xúc và pinned. Nội dung gần trùng được hợp nhất. Không gửi toàn bộ lịch sử dài vào model.")
    heading(doc, "8.1 Background pipeline", 2)
    code_block(doc, "User message + character replies → GPT-5-nano structured analysis → confidence gate → mood observation / relationship event / reflection / durable memory")
    bullet(doc, "Chạy sau khi response đã được gửi nên không chặn đường trả lời chính.")
    bullet(doc, "Thất bại được nuốt an toàn; chat vẫn hoạt động.")
    bullet(doc, "Chỉ lưu mood khi confidence ≥ 0,60; memory yêu cầu confidence ≥ 0,68 và có ý nghĩa.")
    bullet(doc, "Reflection riêng tư không được đọc nguyên văn cho user và không được đổi identity lõi.")

    heading(doc, "9. Tài liệu, tri thức và truy xuất", 1)
    heading(doc, "9.1 Hai workflow tách biệt", 2)
    table(doc, ["Workflow", "Mục đích", "Vòng đời"], [
        ("Room knowledge", "Làm nguồn kiến thức để AI trả lời trong đúng phòng chat", "Tồn tại cùng phòng; xóa phòng thì xóa tài liệu/chunks"),
        ("Character source", "Dùng cốt truyện/tài liệu để biên dịch personality", "Sau khi tạo xong, dữ kiện cần thiết vào CharacterSpec và nguồn tạm bị xóa"),
    ], [4, 6, 6])
    heading(doc, "9.2 Xử lý file", 2)
    bullet(doc, "Định dạng: TXT, MD, PDF, DOCX, PPTX, CSV, XLSX; giới hạn 10 MB.")
    bullet(doc, "Tách section/chunk khoảng 750 từ, overlap 80 từ và lưu metadata trang/chương khi có.")
    bullet(doc, "LLM tạo overview ngắn để tăng khả năng truy xuất chủ đề tổng quát.")
    bullet(doc, "Retrieval dùng cosine similarity trên bag-of-words; khi user nhắc rõ tài liệu/file, hệ thống cho phép đưa top chunks ngay cả khi score thấp.")
    bullet(doc, "Link ngoài kiểm tra scheme, credential và IP global để giảm SSRF; redirect cũng được kiểm tra lại.")
    doc.add_paragraph("Lưu ý kiến trúc: đây là RAG mức MVP nhưng chưa dùng embedding model/pgvector dù .env có cấu hình dự kiến. Khi scale cần thay retrieval lexical bằng hybrid vector + keyword + reranker.")

    heading(doc, "10. Chat 1–1, chat nhóm và streaming", 1)
    heading(doc, "10.1 SSE event flow", 2)
    code_block(doc, """POST /messages/stream
  → user_message     (FE hiển thị ngay)
  → typing           (nhân vật đang soạn)
  → character_message (từng nhân vật xong là phát ngay)
  → character_error  (nếu có)
  → done""")
    bullet(doc, "Chat 1–1 và nhóm dùng chung cơ chế phản hồi tức thời.")
    bullet(doc, "Trong nhóm, route chọn người trả lời; tag duy nhất chỉ gọi nhân vật được tag.")
    bullet(doc, "Mỗi AI chỉ nhận user history và reply của chính mình; reply AI khác không đi vào prompt.")
    bullet(doc, "Điều này tránh lây giọng, tránh AI trả lời theo nội dung của AI khác và giữ independence.")

    heading(doc, "11. Suggestion Engine", 1)
    doc.add_paragraph("Suggestion được sinh từ góc nhìn người dùng, cùng ngôn ngữ và đại từ, gồm đúng ba câu tiếp theo ngắn. Answer và suggestion chạy song song bằng ThreadPoolExecutor để giảm latency tổng. UI chỉ hiện suggestion của reply mới nhất; click sẽ gửi ngay như tin nhắn user.")
    bullet(doc, "Format nội bộ: @@suggest 1|suggest 2|suggest 3.")
    bullet(doc, "Parser loại trùng và validate đúng ba mục; có fallback khi LLM lỗi.")
    bullet(doc, "Ba hướng: đào sâu chi tiết, yêu cầu ý kiến/hành động, và hướng cảm xúc/sáng tạo.")

    heading(doc, "12. Kiểm duyệt và an toàn", 1)
    code_block(doc, "Keyword screening → suspicious excerpts → omni-moderation-latest (nếu cấu hình) → GPT-5-nano contextual classification → allow / block / review")
    table(doc, ["Lớp", "Mục đích", "Hành vi lỗi"], [
        ("Keyword", "Lọc nhanh nhóm sexual, violence, self-harm, hate/extremism, drugs", "Không có network latency"),
        ("OpenAI Moderation", "Tín hiệu phân loại chuyên dụng", "Không khả dụng thì chuyển lớp tiếp theo"),
        ("Contextual LLM", "Phân biệt nhắc tới hợp lệ với hướng dẫn gây hại", "Review/block theo fail-safe"),
        ("Application boundary", "Consent, privacy, anti-dependency, không nhận là người thật", "Được lặp lại trong prompt"),
    ], [3.5, 8, 4.5])
    doc.add_paragraph("Phạm vi kiểm duyệt hiện ưu tiên input và tài liệu do user tạo. Đây chưa thay thế trust & safety policy hoàn chỉnh, human review, audit log bất biến hoặc hệ thống kháng prompt injection chuyên sâu.")

    heading(doc, "13. Mô hình dữ liệu", 1)
    table(doc, ["Bảng", "Nội dung"], [
        ("users", "User demo"), ("characters", "CharacterSpec hiện hành"), ("character_versions", "Phiên bản nhân vật"),
        ("rooms / room_members", "Phòng và thành viên"), ("messages", "Tin nhắn user/AI/document"),
        ("message_alternatives", "Các câu trả lời regenerate"), ("message_suggestions", "Ba gợi ý theo response"),
        ("documents / document_chunks", "Metadata và chunks"), ("room_documents", "Tài liệu gắn theo phòng"),
        ("character_knowledge", "Liên kết tri thức nhân vật"), ("memories", "Ký ức theo visibility"),
        ("relationships", "Familiarity/trust/affinity/tension/respect"), ("character_states", "Emotion/mood/intensity/energy"),
        ("relationship_events", "Sự kiện quan hệ có confidence"), ("mood_observations", "Suy luận cảm xúc user"),
        ("character_reflections", "Reflection riêng tư"), ("prompt_traces", "Lớp context đã ghép"),
        ("daily_tasks", "Công việc hằng ngày"), ("style_examples", "Ví dụ giọng nói bổ sung"),
    ], [5, 11])
    doc.add_paragraph("Database hiện là SQLite tại /data/character_chat.sqlite trên Modal Volume character-chat-data. Mô hình này phù hợp demo/MVP một instance nhưng không phù hợp concurrent write lớn.")

    heading(doc, "14. API", 1)
    table(doc, ["Nhóm", "Endpoint tiêu biểu", "Chức năng"], [
        ("System", "GET /health, /docs", "Health/build version và Swagger có Basic Auth"),
        ("Preset", "GET /v1/presets", "Danh sách template"),
        ("Character", "GET/POST/PATCH/DELETE /v1/characters", "CRUD và đồng bộ preset"),
        ("Generate", "/generate, /mirror, /from-document", "Tạo bằng LLM, bản sao AI và từ truyện"),
        ("Room", "GET/POST/PATCH/DELETE /v1/rooms", "Quản lý phòng"),
        ("Chat", "/messages, /messages/stream, /regenerate", "Chat thường, SSE và trả lời lại"),
        ("Document", "/documents, /knowledge, /character-sources", "Upload, trạng thái, xóa và tách workflow"),
        ("Memory", "/memories", "Xem, pin/sửa và xóa memory"),
        ("Relationship", "/relationship", "State quan hệ và cảm xúc"),
        ("Task", "/daily-tasks", "CRUD task hỗ trợ đời sống"),
        ("Debug", "/prompt-traces", "Quan sát context layer, memory và document IDs"),
    ], [3, 6, 7])
    doc.add_paragraph("Ownership hiện dựa trên header X-Demo-User-Id. Đây chỉ là cơ chế phân vùng demo, không phải xác thực danh tính.")

    heading(doc, "15. Frontend", 1)
    bullet(doc, "React 19 + TypeScript; Vite chịu trách nhiệm dev/build.")
    bullet(doc, "Một SPA phục vụ từ FastAPI; asset có no-cache để production nhận bản build mới sau deploy.")
    bullet(doc, "UI gồm chọn vai trò, preset giải trí, nhân vật của bạn, lịch sử phòng, tạo nhân vật, chat, upload và suggestion.")
    bullet(doc, "Optimistic display cho user message và document card; typing dots tồn tại tới khi response về.")
    bullet(doc, "Xóa character, room và tài liệu; lỗi upload trong chat được hiển thị gọn là ‘Lỗi!’. ")
    doc.add_paragraph("Giới hạn kỹ thuật: state và view hiện tập trung trong main.tsx, phù hợp MVP nhưng nên tách component, hooks, API client và state store khi mở rộng.")

    heading(doc, "16. LLM, fallback và chi phí", 1)
    table(doc, ["Tác vụ", "Model", "Cơ chế"], [
        ("Chat chính", "openai/gpt-4.1-mini qua AVIBI", "Temperature theo verbosity; giới hạn token theo nhân vật"),
        ("Fallback", "zenlayer/gpt-5-nano", "Dùng khi timeout/network hoặc HTTP 408/429/5xx"),
        ("Background analysis", "zenlayer/gpt-5-nano", "Structured JSON sau response"),
        ("Context moderation", "zenlayer/gpt-5-nano", "Chỉ gọi khi keyword đáng ngờ"),
        ("OpenAI moderation", "omni-moderation-latest", "Tùy chọn qua key riêng"),
    ], [4, 5, 7])
    bullet(doc, "Không có request LLM thì không tiêu token AVIBI; lưu dữ liệu và mở trang không tự phát sinh chat token.")
    bullet(doc, "Một lượt thông thường có answer + suggestion song song, sau đó có thể thêm một background analysis.")
    bullet(doc, "Chi phí cần đo thực tế theo input context, output token, tỷ lệ cache/fallback và số nhân vật trả lời trong group.")

    heading(doc, "17. Bảo mật", 1)
    table(doc, ["Hạng mục", "Đã có", "Cần nâng cấp"], [
        ("Secret", "Modal Secret; vault cục bộ mã hóa", "Rotation, RBAC và audit định kỳ"),
        ("API docs", "HTTP Basic", "SSO/VPN hoặc admin gateway"),
        ("Data ownership", "Kiểm tra user_id cho resource", "JWT/OAuth và tenant boundary ở DB"),
        ("Upload", "Extension, 10 MB, parsing, moderation", "Malware scan, MIME magic, sandbox parser"),
        ("URL ingestion", "Chặn credential/private IP và recheck redirect", "Egress proxy và domain policy"),
        ("Prompt", "Voice lock và safety instruction", "Prompt injection filters và output policy"),
        ("CORS", "Cấu hình qua biến môi trường", "Không dùng wildcard ở production SaaS"),
    ], [3.5, 6, 6.5])
    doc.add_paragraph("Không đưa secret thật vào source, README hoặc báo cáo. secrets.vault.json chỉ được dùng cục bộ và phải nằm trong .gitignore; mất master password thì không thể giải mã vault.")

    heading(doc, "18. Deployment và vận hành", 1)
    heading(doc, "18.1 Modal runtime", 2)
    bullet(doc, "Python 3.12 Debian slim; FastAPI chạy bằng @modal.asgi_app().")
    bullet(doc, "0,25 CPU, timeout 600 giây, scaledown 60 giây, tối đa một container.")
    bullet(doc, "Volume /data giữ SQLite qua deployment.")
    bullet(doc, "Ba Modal secrets: character-chat-llm, character-chat-moderation, character-chat-docs.")
    bullet(doc, "Frontend được build trước deploy và mount dist vào runtime image.")
    heading(doc, "18.2 Lệnh vận hành chính", 2)
    code_block(doc, """# Kiểm thử
python -m pytest -q
cd frontend
npm run build

# Deploy production trên PowerShell UTF-8
$env:PYTHONUTF8='1'
modal deploy modal_app.py --stream-logs

# Xem log
modal app logs character-chat -f

# Health
Invoke-RestMethod -Uri "https://nguyenthanhcongloi332004--character-chat-web.modal.run/health" -TimeoutSec 30""")
    doc.add_paragraph("Các lệnh secret, vault, backup và troubleshooting chi tiết được duy trì trong README.md của repository.")

    heading(doc, "19. Kiểm thử", 1)
    table(doc, ["Lớp", "Nội dung", "Trạng thái gần nhất"], [
        ("Unit/domain", "Preset, prompt, relationship, state, routing, suggestion", "11 test passed"),
        ("API", "Character, room, message, documents, relationship, trace", "Nằm trong tests/test_app.py"),
        ("Frontend", "TypeScript noEmit + Vite production build", "Build thành công"),
        ("Smoke", "Modal endpoints và deployed assets", "Có script riêng"),
        ("UX timing", "Thứ tự user message/typing/reply", "Có Playwright-style script ui_timing.mjs"),
    ], [3, 8, 5])
    doc.add_paragraph("Khoảng trống: chưa có load test, chaos test, security test tự động, eval định lượng về character consistency, hallucination và moderation precision/recall.")

    heading(doc, "20. Đánh giá hiện trạng và rủi ro", 1)
    table(doc, ["Rủi ro", "Mức", "Ảnh hưởng", "Khuyến nghị"], [
        ("Demo identity bằng header", "Cao", "Không đủ cho dữ liệu người dùng thật", "Triển khai OAuth/JWT và tenant authorization"),
        ("SQLite một container", "Cao", "Giới hạn scale và concurrent write", "PostgreSQL managed + migration"),
        ("Lexical retrieval", "Trung bình", "Bỏ sót ngữ nghĩa và từ đồng nghĩa", "Embedding + pgvector + hybrid rerank"),
        ("LLM external dependency", "Trung bình", "Latency, outage và chi phí", "Timeout budget, circuit breaker, metrics và multi-provider"),
        ("Character quality chưa có eval", "Trung bình", "Khó biết personality drift", "Golden conversations và automated judge"),
        ("Moderation policy đơn giản", "Trung bình", "False positive/negative", "Policy matrix, logging, review và appeal"),
        ("Monolithic frontend", "Thấp/Trung bình", "Khó bảo trì khi thêm chức năng", "Tách modules và typed API client"),
    ], [4, 2, 5, 5])

    heading(doc, "21. Lộ trình SaaS đề xuất", 1)
    table(doc, ["Giai đoạn", "Mục tiêu", "Đầu việc"], [
        ("P0 – Cứng hóa MVP", "Dùng thử có kiểm soát", "Auth thật, Postgres, migrations, error telemetry, rate limit, backup/restore"),
        ("P1 – Chất lượng AI", "Tăng retention", "Character eval, semantic memory, pgvector, prompt versioning, A/B testing"),
        ("P2 – Thương mại", "Monetization", "Subscription, token quota, billing, usage dashboard, cost controls"),
        ("P3 – Mở rộng trải nghiệm", "Tăng engagement", "Avatar expression, push reminders có consent, mobile/PWA; TTS nếu có nhu cầu"),
        ("P4 – Scale", "Nhiều tenant", "Worker queue, async ingestion, caching, autoscaling, observability và SLO"),
    ], [3, 4, 9])
    heading(doc, "21.1 KPI nên theo dõi", 2)
    bullet(doc, "Activation: tỷ lệ tạo nhân vật/phòng và gửi tin nhắn đầu tiên.")
    bullet(doc, "Engagement: messages/session, session length, suggestion CTR, upload usage.")
    bullet(doc, "Retention: D1/D7/D30 theo loại nhân vật.")
    bullet(doc, "Quality: role consistency, memory precision, grounded-document answer rate, regeneration rate.")
    bullet(doc, "Reliability: p50/p95 latency, LLM error/fallback rate, ingestion failure, moderation latency.")
    bullet(doc, "Economics: token/user/day, cost per retained user, gross margin theo plan.")

    heading(doc, "22. Hướng dẫn bàn giao", 1)
    numbered(doc, "Đọc README.md để nắm lệnh deploy, secret, vault, backup và xử lý lỗi.")
    numbered(doc, "Đọc presets.py và prompts.py để hiểu tính cách cùng quy tắc prompt.")
    numbered(doc, "Đọc engine.py theo thứ tự: LLM client → route → state/relationship → assemble.")
    numbered(doc, "Đọc main.py theo nhóm endpoint Character → Room → Chat → Document → Memory.")
    numbered(doc, "Chạy python -m pytest -q và npm run build trước mọi deployment.")
    numbered(doc, "Kiểm tra /health, /v1/presets và một lượt SSE sau deployment.")
    numbered(doc, "Không commit secret thật; thay key bằng setup script hoặc Modal Secret CLI.")

    doc.add_page_break()
    heading(doc, "Phụ lục A – Workflow đầy đủ của một lượt chat", 1)
    code_block(doc, """1. FE gửi message qua SSE endpoint.
2. Backend kiểm duyệt input.
3. Lưu user message và phát event user_message ngay.
4. Router chọn nhân vật dựa trên tag và speaker mode.
5. Với từng nhân vật: phát typing.
6. Load CharacterSpec, relationship, state, recent history, memory, reflection, task và document chunks.
7. Prompt Compiler ghép context theo lớp và ghi PromptTrace.
8. Chạy answer và suggestions song song.
9. Validate suggestion; dùng fallback khi cần.
10. Lưu character message, sources và suggestions.
11. Cập nhật relationship chậm và phát character_message.
12. Lặp lại cho nhân vật tiếp theo trong group.
13. Trích memory theo rule nhanh.
14. Phát done.
15. Background Nano phân tích mood, event, reflection và durable memory.""")

    heading(doc, "Phụ lục B – Danh sách endpoint", 1)
    endpoints = [
        "GET /health", "GET /docs", "GET /openapi.json", "GET /v1/presets",
        "GET|POST /v1/characters", "GET|PATCH|DELETE /v1/characters/{id}",
        "POST /v1/characters/generate", "POST /v1/characters/mirror", "POST /v1/characters/from-document",
        "GET|POST /v1/rooms", "GET|PATCH|DELETE /v1/rooms/{id}",
        "GET|POST /v1/rooms/{id}/messages", "POST /v1/rooms/{id}/messages/stream",
        "POST /v1/rooms/{id}/messages/{message_id}/regenerate", "POST /v1/rooms/{id}/knowledge",
        "GET|POST /v1/documents", "POST /v1/documents/from-url", "GET|DELETE /v1/documents/{id}",
        "POST /v1/documents/{id}/discover-characters", "POST /v1/character-sources",
        "GET /v1/characters/{id}/memories", "GET /v1/rooms/{id}/memories", "PATCH|DELETE /v1/memories/{id}",
        "GET /v1/characters/{id}/relationship", "GET|POST /v1/daily-tasks", "PATCH|DELETE /v1/daily-tasks/{id}",
        "GET /v1/debug/rooms/{id}/prompt-traces",
    ]
    for endpoint in endpoints:
        bullet(doc, endpoint)

    heading(doc, "Phụ lục C – Secrets và dữ liệu vận hành", 1)
    table(doc, ["Tên", "Nội dung", "Nơi lưu"], [
        ("character-chat-llm", "AIMP_API_KEY", "Modal Secret"),
        ("character-chat-moderation", "OPENAI_MODERATION_API_KEY", "Modal Secret"),
        ("character-chat-docs", "DOCS_USERNAME, DOCS_PASSWORD", "Modal Secret"),
        ("character-chat-data", "character_chat.sqlite", "Modal Volume /data"),
        ("secrets.vault.json", "Bản sao secret mã hóa tùy người vận hành", "Máy local, không commit"),
    ], [4.5, 5.5, 6])

    heading(doc, "Phụ lục D – Nguồn tham khảo kỹ thuật", 1)
    doc.add_paragraph("Thiết kế đã tham khảo các mô hình phổ biến trong TavernAI, open-character-ai, SillyTavern, RisuAI, Shikigami Protocol và her-go ở cấp ý tưởng: character card, prompt assembly, lore/context retrieval, group chat, emotion state, mood agent, reflection và memory. Mã dự án hiện là triển khai riêng. Việc tái sử dụng trực tiếp mã nguồn bên thứ ba phải tuân thủ giấy phép; xem THIRD_PARTY_NOTICES.md.")

    doc.add_paragraph()
    p = doc.add_paragraph("HẾT BÁO CÁO")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.color.rgb = RGBColor.from_string(PURPLE)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build()
    print(path)
