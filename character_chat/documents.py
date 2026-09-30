import csv, io, json, re
from pathlib import Path
from .retrieval import chunks
from .schemas import DiscoveryCandidate, uid
from .db import DocumentChunk

def parse_document(filename, data):
    ext = Path(filename).suffix.lower()
    sections = []
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise ValueError("Invalid PDF file format.")
    if ext in (".docx", ".pptx", ".xlsx") and not data.startswith(bytes((0x50, 0x4b, 0x03, 0x04))):
        raise ValueError("Invalid Office file format.")
    if ext in (".txt", ".md"):
        text = data.decode("utf-8-sig", errors="replace")
        pattern = r"(?im)(?=^\s*(?:chapter|ch\u01b0\u01a1ng)\s+\d+)"
        for n, part in enumerate(re.split(pattern, text)):
            if part.strip():
                match = re.search(r"(?im)^\s*(?:chapter|ch\u01b0\u01a1ng)\s+(\d+)", part)
                sections.append((part, {"chapter": int(match.group(1)) if match else n + 1, "section": n + 1}))
    elif ext == ".pdf":
        import fitz
        with fitz.open(stream=data, filetype="pdf") as pdf:
            if len(pdf) > 500:
                raise ValueError("PDF exceeds the 500 page processing limit.")
            for i, page in enumerate(pdf):
                sections.append((page.get_text(), {"page": i + 1}))
    elif ext == ".docx":
        from docx import Document
        doc = Document(io.BytesIO(data)); chapter = 0
        for paragraph in doc.paragraphs:
            if paragraph.style.name.lower().startswith("heading"):
                chapter += 1
            if paragraph.text.strip():
                sections.append((paragraph.text, {"chapter": max(1, chapter), "section": paragraph.style.name}))
        for table_index, table in enumerate(doc.tables, 1):
            for row_index, row in enumerate(table.rows, 1):
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    sections.append((" | ".join(cells), {"section": f"table {table_index}", "row": row_index}))
    elif ext == ".pptx":
        from pptx import Presentation
        for i, slide in enumerate(Presentation(io.BytesIO(data)).slides):
            sections.append(("\n".join(shape.text for shape in slide.shapes if shape.has_text_frame), {"slide": i + 1}))
    elif ext == ".csv":
        rows = csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace")))
        for i, row in enumerate(rows):
            sections.append((" | ".join(row), {"row": i + 1}))
    elif ext == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in workbook:
            for i, row in enumerate(sheet.values):
                sections.append((" | ".join(str(value or "") for value in row), {"sheet": sheet.title, "row": i + 1}))
    else:
        raise ValueError("Unsupported file type")
    extracted = [(text, meta) for text, meta in sections if text and text.strip()]
    if not extracted:
        raise ValueError("No readable text found. Scanned PDFs need OCR before upload.")
    return extracted

def save_chunks(db, document_id, sections):
    for text, meta in sections:
        for part, chunk_meta in chunks(text, meta):
            db.add(DocumentChunk(id=uid(), document_id=document_id, text=part, meta=json.dumps(chunk_meta)))
    db.commit()

def discover_fallback(sections):
    text = "\n".join(value for value, _ in sections)[:12000]
    names = {}
    for match in re.finditer(r"(?<!\w)([A-Z\u00c0-\u1ef8][a-z\u00e0-\u1ef9]{2,})(?=\s*(?:n\u00f3i|said|h\u1ecfi|asked|:))", text):
        name = match.group(1); names[name] = names.get(name, 0) + 1
    candidates = []
    for name, count in sorted(names.items(), key=lambda item: -item[1])[:10]:
        match = re.search(r".{0,80}\b" + re.escape(name) + r"\b.{0,100}", text)
        candidates.append(DiscoveryCandidate(name=name, description="Nh\u00e2n v\u1eadt " + name + " xu\u1ea5t hi\u1ec7n trong t\u00e0i li\u1ec7u", importance=min(1, count / 10), confidence=.5, source_evidence=[match.group(0) if match else name], dialogue_count=count))
    return candidates
